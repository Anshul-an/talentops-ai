from openpyxl import load_workbook
from io import BytesIO

from app import db
from app.excel_export import build_workbook
from app.scorer import heuristic_score


def test_workbook_has_three_required_sheets_and_pipeline_counts():
    cid, _ = db.upsert_candidate(
        {
            "full_name": "Ananya Rao",
            "email": "xlsx@example.com",
            "source_channel": "Naukri",
            "skills": "Python, FastAPI, pytest",
            "current_ctc": 24,
            "expected_ctc": 32,
        }
    )
    jid = db.insert_job("Backend", "Meridian", "Python FastAPI pytest Claude")
    cand = db.get_candidate(cid)
    job = db.get_job(jid)
    scored = heuristic_score(cand, job, 70)
    db.insert_score({**scored, "candidate_id": cid, "job_id": jid})

    data = build_workbook()
    wb = load_workbook(BytesIO(data))
    assert "Master Daily Pipeline" in wb.sheetnames
    assert "Recruiter KPI Leaderboard" in wb.sheetnames
    assert "Channel Performance Matrix" in wb.sheetnames

    pipeline = wb["Master Daily Pipeline"]
    headers = [c.value for c in next(pipeline.iter_rows(min_row=1, max_row=1))]
    assert "Candidate Name" in headers
    assert "Sourcing Channel" in headers
    names = [row[0].value for row in pipeline.iter_rows(min_row=2) if row[0].value]
    assert "Ananya Rao" in names

    kpi = wb["Recruiter KPI Leaderboard"]
    kpi_headers = [c.value for c in next(kpi.iter_rows(min_row=1, max_row=1))]
    assert "Login Hours (Kredily)" in kpi_headers
    assert "Productive Hours (TeamLogger)" in kpi_headers
